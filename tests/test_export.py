






"""Excel 导出接口测试。

覆盖 6 个导出端点：
- 库存流水、采购订单、销售订单、低库存预警、采购统计、销售统计

断言：200 + Content-Type + 文件名 + openpyxl 可打开。
"""

from __future__ import annotations

import io

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from models.role import Permission, Role, RolePermission


# 需要的查看类权限码
VIEW_PERMS = [
    "inventory:view",
    "purchase:view",
    "sales:view",
    "reports:view",
]


@pytest.fixture
def setup_operator_perms(db_session: Session):
    """创建 operator 角色并分配 4 个查看类权限。

    auth_client 创建的用户 role_id 为 None，默认按 operator 处理，
    因此需要在测试库中预置 operator 角色及其权限。
    """
    role = Role(name="operator", description="普通操作员")
    db_session.add(role)
    db_session.flush()

    perms = [
        Permission(code=code, name=f"查看{code.split(':')[0]}", module=code.split(":")[0])
        for code in VIEW_PERMS
    ]
    db_session.add_all(perms)
    db_session.flush()

    for p in perms:
        db_session.add(RolePermission(role_id=role.id, permission_id=p.id))
    db_session.commit()
    return role


def _assert_xlsx(resp):
    """断言响应是合法的 xlsx 文件。"""
    assert resp.status_code == 200
    assert "spreadsheetml.sheet" in resp.headers["content-type"]
    cd = resp.headers["content-disposition"]
    assert "attachment" in cd
    assert ".xlsx" in cd

    wb = load_workbook(io.BytesIO(resp.content))
    assert len(wb.sheetnames) >= 1
    wb.close()


@pytest.mark.asyncio
async def test_export_inventory_movements(auth_client: AsyncClient, setup_operator_perms):
    resp = await auth_client.get("/api/v1/inventory/movements/export")
    _assert_xlsx(resp)


@pytest.mark.asyncio
async def test_export_purchase_orders(auth_client: AsyncClient, setup_operator_perms):
    resp = await auth_client.get("/api/v1/purchases/orders/export")
    _assert_xlsx(resp)


@pytest.mark.asyncio
async def test_export_sales_orders(auth_client: AsyncClient, setup_operator_perms):
    resp = await auth_client.get("/api/v1/sales/orders/export")
    _assert_xlsx(resp)


@pytest.mark.asyncio
async def test_export_low_stock_alerts(auth_client: AsyncClient, setup_operator_perms):
    resp = await auth_client.get("/api/v1/reports/low-stock-alerts/export")
    _assert_xlsx(resp)


@pytest.mark.asyncio
async def test_export_purchase_summary(auth_client: AsyncClient, setup_operator_perms):
    resp = await auth_client.get(
        "/api/v1/reports/purchase-summary/export",
        params={"start_date": "2025-01-01", "end_date": "2025-12-31"},
    )
    _assert_xlsx(resp)


@pytest.mark.asyncio
async def test_export_sales_summary(auth_client: AsyncClient, setup_operator_perms):
    resp = await auth_client.get(
        "/api/v1/reports/sales-summary/export",
        params={"start_date": "2025-01-01", "end_date": "2025-12-31"},
    )
    _assert_xlsx(resp)
